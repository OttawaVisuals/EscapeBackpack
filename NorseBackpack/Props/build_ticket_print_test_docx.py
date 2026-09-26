"""
Applies the postcard PrintTest technique (Postcards/build_print_test_docx.py, decided
25 Sept 2026) to the two tickets parked under PR-17/PR-18 for home-duplex registration
drift: the L'Anse museum ticket (PZ-12) and the Rouen ticket (PZ-15).

Same fix, same reason: alignment on this printer is reliable in the TOP zone of the
page and unreliable near the bottom, so each ticket is rotated 90 deg to landscape --
both are narrower than they are tall, so rotating puts the short dimension (2 in) as
the printed height, well inside the reliable zone. Page 1 = front, page 2 = back,
rotated the extra 180 deg for a top-to-bottom flip.

Untested assumption carried over from the postcard fix: the 1.5mm rightward shift on
the back page that cancelled drift there is reused here as a starting point, not
re-measured for this stock (176 gsm cream cardstock vs. the postcard stock). Check the
first printed sheet before trusting it -- if the back is still off, adjust TEXT_SHIFT
for these two files independently of the postcard value.

Source: rasterized straight from each ticket's own Print.pdf (already at exact trim
size, 300dpi) via PyMuPDF, not from any intermediate PNG, so this can never drift from
the production build. Rebuild after either build_museum_ticket_pdf.py or
build_rouen_ticket_pdf.py changes.
"""

import io
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Inches, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output" / "docx"
PDF_DIR = ROOT / "output" / "pdf"

RENDER_DPI = 300
TOP_MARGIN_IN = 0.3
SIDE_MARGIN_IN = 0.6
ROTATE_DEGREES = -90  # clockwise, same convention as the postcard script
TEXT_ROTATE_DEGREES = ROTATE_DEGREES + 180
# Carried over from the postcard fix (PR-23) as a starting point -- not yet re-measured
# for this ticket stock. Verify on the first printed sheet.
TEXT_SHIFT = Mm(1.5)

TICKETS = {
    "MuseumTicket": {
        "pdf": PDF_DIR / "Museum_Ticket_Print.pdf",
        "w_in": 2.0,
        "h_in": 5.5,
    },
    "RouenTicket": {
        "pdf": PDF_DIR / "Rouen_Ticket_Print.pdf",
        "w_in": 2.0,
        "h_in": 3.0,
    },
}


def page_to_image_stream(pdf_path, page_index, rotate_degrees):
    doc = fitz.open(pdf_path)
    page = doc[page_index]
    zoom = RENDER_DPI / 72
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    rotated = img.rotate(rotate_degrees, expand=True)
    buf = io.BytesIO()
    rotated.save(buf, format="PNG")
    buf.seek(0)
    return buf


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {qn("w:val"): "none"})
        borders.append(el)
    tbl_pr.append(borders)


def build_page(doc, pdf_path, page_index, w_in, h_in, rotate_degrees):
    # Rotated 90 deg: printed width/height swap.
    printed_w_in, printed_h_in = h_in, w_in
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    row.height = Inches(printed_h_in)
    cell = row.cells[0]
    cell.width = Inches(printed_w_in)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run()
    run.add_picture(
        page_to_image_stream(pdf_path, page_index, rotate_degrees),
        width=Inches(printed_w_in),
        height=Inches(printed_h_in),
    )


def set_margins(section, horizontal_shift=0):
    section.top_margin = Inches(TOP_MARGIN_IN)
    section.left_margin = Inches(SIDE_MARGIN_IN) + horizontal_shift
    section.right_margin = Inches(SIDE_MARGIN_IN) - horizontal_shift
    section.bottom_margin = Inches(0.3)


def build_ticket(name, spec):
    doc = Document()
    set_margins(doc.sections[0])
    build_page(doc, spec["pdf"], 0, spec["w_in"], spec["h_in"], ROTATE_DEGREES)

    text_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(text_section, horizontal_shift=TEXT_SHIFT)
    build_page(doc, spec["pdf"], 1, spec["w_in"], spec["h_in"], TEXT_ROTATE_DEGREES)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / f"PrintTest_{name}.docx"
    doc.save(str(out_path))
    return out_path


def main():
    for name, spec in TICKETS.items():
        print(build_ticket(name, spec))


if __name__ == "__main__":
    main()
