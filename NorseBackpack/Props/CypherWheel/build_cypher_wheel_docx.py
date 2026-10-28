"""Word version of NorseBackpack/Print/_pdf_source/Cypher_Wheel_Print.pdf (5 Letter pages), for printers that handle the PDF badly.

Each page is rasterised at 300 dpi from the PDF (never drifts from the builder). The separate pieces on a page
(for example wheel III and the arm) are cut apart at the blank bands between them and placed again at exact
size, closer together and centred vertically, so the blank space sits at the top and bottom edges. Pieces
keep their left-right position and size; only the vertical gaps change. Rebuild after build_cypher_wheel_pdf.py.
Run from the repo root:  python NorseBackpack/Props/CypherWheel/build_cypher_wheel_docx.py
"""
import importlib.util
import io
from pathlib import Path

import numpy as np
import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[3]
_spec = importlib.util.spec_from_file_location(
    "postcard_print", ROOT / "NorseBackpack" / "Postcards" / "build_print_test_docx.py")
_pc = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_pc)

DPI = 300
PAD_IN = 0.05          # clear space kept around each piece when it is cut out
GAP_IN = 0.45          # space between pieces on the same page
# Blank bands narrower than this keep the pieces on either side together (page 3's arm and its caption).
MERGE_IN = {1: 0.3, 2: 0.3, 3: 0.3, 4: 9, 5: 9}   # pages 4-5: one block, only centred


def pieces(page, merge_in):
    pix = page.get_pixmap(dpi=100, alpha=False)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)[:, :, :3]
    occ = (a < 245).any(axis=2).any(axis=1)
    bands, start = [], None
    for y, o in enumerate(occ):
        if o and start is None:
            start = y
        if not o and start is not None:
            bands.append([start, y])
            start = None
    if start is not None:
        bands.append([start, len(occ)])
    merged = []
    for b in bands:
        if merged and (b[0] - merged[-1][1]) / 100 < merge_in:
            merged[-1][1] = b[1]
        else:
            merged.append(b)
    return [(max(0.0, y0 / 100 - PAD_IN), min(page.rect.height / 72, y1 / 100 + PAD_IN)) for y0, y1 in merged]


doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(s, side, Inches(0))
pdf = pymupdf.open(ROOT / "NorseBackpack" / "Print" / "_pdf_source" / "Cypher_Wheel_Print.pdf")
for n, page in enumerate(pdf, 1):
    w = page.rect.width / 72
    cuts = pieces(page, MERGE_IN[n])
    total = sum(y1 - y0 for y0, y1 in cuts) + GAP_IN * (len(cuts) - 1)
    y = (11 - total) / 2
    para = doc.add_paragraph()
    if n > 1:
        doc.paragraphs[-2].add_run().add_break(WD_BREAK.PAGE)
    for y0, y1 in cuts:
        clip = pymupdf.Rect(0, y0 * 72, page.rect.width, y1 * 72)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(DPI / 72, DPI / 72), clip=clip, alpha=False)
        _pc.add_floating_picture(para, io.BytesIO(pix.tobytes("png")), w, y1 - y0, 0, y)
        y += (y1 - y0) + GAP_IN
    print(n, [(round(a, 2), round(b, 2)) for a, b in cuts], "top margin", round((11 - total) / 2, 2), "in")
out = ROOT / "NorseBackpack" / "Print" / "Print_CypherWheel.docx"
doc.save(str(out))
print(out)
