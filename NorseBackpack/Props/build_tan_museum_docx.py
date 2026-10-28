# PRINT EXPORT SUPERSEDED 20 Oct 2026: use Props/build_museum_bookmarks.py for current 2 x 6 tickets.
"""Background-free tan-stock variants of the adopted museum ticket PDFs.

Preserves PDF text and artwork; removes paper-colour vector fills and applies
a PDF colour-key mask to the Brattahlid illustration's pale cream ground.
Does not regenerate artwork or run the superseded layout builders.
"""
from pathlib import Path
import importlib.util
import io
import json

from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, NameObject, NumberObject, FloatObject, TextStringObject, ByteStringObject
import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
PRINT = ROOT / 'Print'
WORK = ROOT.parents[1] / 'output/tan-tickets'
WORK.mkdir(parents=True, exist_ok=True)
spec = importlib.util.spec_from_file_location('print_positions', ROOT / 'Postcards/build_print_test_docx.py')
positions = importlib.util.module_from_spec(spec)
spec.loader.exec_module(positions)
PAPER = (.960784, .937255, .874510)
PALE = (.913725, .901961, .850980)
INK = (.156863, .231373, .203922)


def matches(a, b):
    return len(a) == len(b) and all(abs(float(x)-y) < .00001 for x, y in zip(a, b))


def edit_pdf(source, out, museum):
    reader = PdfReader(source)
    writer = PdfWriter()
    summary = []
    for i in range(2):  # Aud's third page is the designer answer copy.
        page = writer.add_page(reader.pages[i])
        original_text = reader.pages[i].extract_text()
        stream = page.get_contents()
        ops = stream.operations
        def preserve_bytes(value):
            if isinstance(value, TextStringObject):
                return ByteStringObject(value.original_bytes)
            if isinstance(value, ArrayObject):
                return ArrayObject([preserve_bytes(x) for x in value])
            return value
        # Preserve subset-font character codes when rewriting the PDF stream.
        ops = [([preserve_bytes(x) for x in args], op) for args, op in ops]
        fill, rectangle, removed = None, None, 0
        for n, (args, op) in enumerate(ops):
            if op == b'rg':
                fill = tuple(float(x) for x in args)
            elif op == b'n':
                rectangle = None
            elif op == b're':
                rectangle = tuple(float(x) for x in args)
            elif op in (b'f', b'f*'):
                paper_rect = rectangle is not None and (matches(fill or (), PAPER) or matches(fill or (), PALE))
                header = museum and i == 0 and rectangle == (0., 336., 252., 96.)
                if paper_rect or header:
                    ops[n] = ([], b'n')
                    removed += 1
                rectangle = None
        if museum and i == 0:
            for n, (args, op) in enumerate(ops):
                if op == b'rg' and matches(args, PAPER):
                    ops[n] = ([FloatObject(x) for x in INK], op)
            images = page['/Resources']['/XObject']
            for ref in images.values():
                image = ref.get_object()
                if image.get('/Subtype') == '/Image':
                    assert image['/ColorSpace'] == '/DeviceRGB'
                    # Pale warm ground only; retain darker paths, numbered markers,
                    # coastline, water, buildings, trees and all underlying bytes.
                    image[NameObject('/Mask')] = ArrayObject([NumberObject(x) for x in (240,255,235,255,215,240)])
        stream.operations = ops
        page.replace_contents(stream)
        assert page.extract_text() == original_text, (i, repr(original_text), repr(page.extract_text()))
        summary.append({'page': i+1, 'removed_background_rectangles': removed,
                        'size_points': [float(page.mediabox.width), float(page.mediabox.height)],
                        'text_preserved': True})
    writer.write(out)
    return summary


def build(name, source, museum):
    clean_pdf = WORK / (name + '_Tan_Source.pdf')
    summary = edit_pdf(source, clean_pdf, museum)
    pdf = pdfium.PdfDocument(str(clean_pdf))
    doc = Document()
    for i in range(2):
        if i:
            doc.add_section(WD_SECTION.NEW_PAGE)
        sec = doc.sections[i]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(.5)
        sec.top_margin = sec.bottom_margin = Inches(.5)
        page = pdf[i]
        w, h = page.get_size()
        assert (w, h) == (252, 432)
        raster = page.render(scale=300/72, fill_color=(255,255,255,0)).to_pil()
        angle = positions.ROTATE_DEGREES if i == 0 else positions.TEXT_ROTATE_DEGREES
        raster = raster.rotate(angle, expand=True)
        buf = io.BytesIO()
        raster.save(buf, format='PNG')
        buf.seek(0)
        # Match the established ticket print method: front at .6 inches from
        # left, reverse mirrored from right; opposite rotation for manual flip.
        x = .6 if i == 0 else 8.5-.6-6
        p = doc.add_paragraph()
        positions.add_floating_picture(p, buf, 6, 3.5, x, .3)
        if i == 0:
            # Hairline trim guide at the actual 6 by 3.5 inch ticket boundary.
            shape_props = p._p.xpath('.//pic:spPr')[0]
            line = OxmlElement('a:ln')
            line.set('w', '3175')  # 0.25 pt
            solid = OxmlElement('a:solidFill')
            color = OxmlElement('a:srgbClr')
            color.set('val', '000000')
            solid.append(color)
            line.append(solid)
            shape_props.append(line)
        raster.save(WORK / f'{name}_face_{i+1}.png')
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title = name.replace('_', ' ') + ' Ticket for Tan Paper'
    doc.core_properties.author = 'The Cryptic Bureau'
    dest = PRINT / (name + '_Tan_Paper.docx')
    doc.save(dest)
    return {'file': str(dest), 'pages': summary}


if __name__ == '__main__':
    results = [build('Brattahlid_Museum_Ticket', PRINT/'_pdf_source/Museum_Ticket_Print.pdf', True),
               build('Aud_Museum_Ticket', PRINT/'_pdf_source/Aud_Ticket_Print.pdf', False)]
    (WORK/'checks.json').write_text(json.dumps(results, indent=2))
    print(json.dumps(results, indent=2))
