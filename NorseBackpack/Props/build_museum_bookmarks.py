"""Four reflowed 2 x 6 inch museum tickets and aligned two-page Letter DOCX proofs.

Puzzle data is read from original data modules without running old layout builders.
The Oslo reverse is clipped/repositioned as PDF content at 100%: never rescaled.
"""
from pathlib import Path
import ast, io, json, importlib.util, shutil
from xml.sax.saxutils import escape
from PIL import Image
from pypdf import PdfReader, PdfWriter, Transformation
from pypdf.generic import RectangleObject, NameObject, ArrayObject, NumberObject
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from docx.oxml import OxmlElement

NORSE=Path(__file__).resolve().parents[1]
ROOT=NORSE.parent
PRINT=NORSE/'Print'; PDF=PRINT/'_pdf_source'
HERE=Path(__file__).resolve().parent
SOURCES=HERE/'MuseumBookmarks/Sources'
WORK=ROOT.parent/'output/museum-2x6';WORK.mkdir(parents=True,exist_ok=True)
W,H=144,432
INK,RUST,RULE=[HexColor(x) for x in ('#283B34','#985326','#B2B5A5')]
for name,file in [('Body','SourceSans-400.ttf'),('Bold','SourceSans-600.ttf'),('Display','Cormorant-600.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(ROOT/'Fonts/CrypticBureau'/file)))
spec=importlib.util.spec_from_file_location('museum_print_positions',NORSE/'Postcards/build_print_test_docx.py')
positions=importlib.util.module_from_spec(spec);spec.loader.exec_module(positions)

def literal(file,name):
    tree=ast.parse(file.read_text(encoding='utf-8'))
    return next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign)
                and any(isinstance(t,ast.Name) and t.id==name for t in n.targets))

INDEX=literal(HERE/'MuseumTicket/build_museum_ticket_pdf.py','INDEX')
LEGS=literal(HERE/'AudTicket/build_aud_ticket_pdf.py','LEGS')
COINS=literal(HERE/'AudTicket/build_aud_ticket_pdf.py','COIN_KEY')
SAXON=literal(HERE/'RouenTicket/build_rouen_ticket_pdf.py','SAXON')
NORMAN=literal(HERE/'RouenTicket/build_rouen_ticket_pdf.py','NORMAN')
RULES=literal(HERE/'Hnefatafl/build_ticket_pdf.py','RULES_LINES')

def text(c,x,y,value,size=10,font='Body',color=INK,align='left'):
    c.setFillColor(color);c.setFont(font,size)
    assert pdfmetrics.stringWidth(value,font,size)<=W-18+.01,(value,size)
    {'left':c.drawString,'center':c.drawCentredString,'right':c.drawRightString}[align](x,y,value)

def rule(c,y,left=10,right=134):
    c.setStrokeColor(RULE);c.setLineWidth(.4);c.line(left,y,right,y)

def base(c):
    # Unprinted substrate for the home prototype, with a small heritage frame.
    c.setStrokeColor(INK);c.setLineWidth(.45)
    c.lines([(0,0,W,0),(W,0,W,H),(W,H,0,H),(0,H,0,0)])
    c.setStrokeColor(RULE);c.setLineWidth(.5);c.rect(5,5,W-10,H-10,fill=0,stroke=1)

def paragraph(c,value,x,top,width,size=11,leading=13,font='Body'):
    lines=[];line=''
    for word in value.split():
        trial=(line+' '+word).strip()
        if pdfmetrics.stringWidth(trial,font,size)<=width:
            line=trial
        else:
            assert line,(word,width)
            lines.append(line);line=word
    if line:lines.append(line)
    for i,line in enumerate(lines):text(c,x,top-i*leading,line,size,font)
    return top-len(lines)*leading

def brattahlid_front(c):
    base(c)
    text(c,72,410,'VIKING MUSEUM',16,'Display',align='center')
    text(c,72,391,'OF BRATTAHLÍÐ',15,'Display',align='center')
    rule(c,379)
    # Reuse the existing tall nine-site illustration at its native 2-inch width.
    # Clip its old baked header and lower serial-number/footer area at placement.
    c.saveState();p=c.beginPath();p.rect(6,65,132,295);c.clipPath(p,stroke=0,fill=0)
    c.drawImage(str(HERE/'MuseumTicket/Museum_Ticket_Front_300dpi_v5_transparent.png'),0,30,144,396,mask='auto')
    c.restoreState()
    rule(c,53)
    text(c,72,36,'GENERAL ADMISSION',9,'Bold',align='center')

def brattahlid_back(c):
    base(c)
    text(c,72,410,'BRATTAHLÍÐ',14,'Display',align='center')
    text(c,72,389,'VISITOR INDEX',16,'Display',align='center');rule(c,378)
    text(c,25,365,'GALLERY',8,'Bold',RUST)
    text(c,132,365,'ROOM',8,'Bold',RUST,'right')
    for i,(letter,name,number) in enumerate(INDEX):
        y=350-i*12.6
        text(c,11,y,letter,10,'Bold',RUST)
        size=min(10,95/pdfmetrics.stringWidth(name,'Body',1))
        assert size>=9.5,(name,size)
        text(c,25,y,name,size)
        text(c,132,y,str(number),10,'Bold',align='right')
    assert y>20

def aud_art():
    page=PdfReader(SOURCES/'Aud_Original_Artwork.pdf').pages[0]
    names=[str(args[0]).lstrip('/') for args,op in page.get_contents().operations if op==b'Do']
    images={im.name.rsplit('.',1)[0]:im.data for im in page.images}
    return [ImageReader(io.BytesIO(images[n])) for n in names]

def aud_front(c):
    base(c)
    text(c,72,410,'AUÐAR SAFN',20,'Display',align='center')
    text(c,72,393,'HVAMMUR · DALIR',9,'Bold',RUST,'center');rule(c,383)
    text(c,72,365,'THE COIN ROOM',15,'Display',align='center')
    text(c,72,351,'What each coin is worth today',9,align='center')
    art=aud_art();image_index=0;y=334
    for group,coins in COINS:
        rule(c,y,10,60);rule(c,y,84,134)
        if group=='fehu':
            c.setStrokeColor(INK);c.setLineWidth(1.2)
            c.line(69,y-7,69,y+7);c.line(69,y+1,76,y+6);c.line(69,y-3,76,y+2)
        else:
            c.drawImage(art[image_index],65,y-7,14,14,mask='auto');image_index+=1
        y-=24
        for name,label,value,rare in coins:
            c.drawImage(art[image_index],12,y-12,25,25,mask='auto');image_index+=1
            text(c,44,y+4,label,10.5)
            text(c,44,y-10,f'${value:,}',11,'Bold',RUST)
            y-=35
        y-=9
    assert image_index==9

def aud_back(c):
    base(c)
    text(c,72,410,'AUÐAR SAFN',14,'Display',align='center')
    text(c,72,388,'TREASURE HUNT',15,'Display',align='center');rule(c,377)
    top=paragraph(c,'Trace the route on your map.',12,362,120,11,13)-8
    for i,leg in enumerate(LEGS):
        text(c,12,top,str(i+1),12,'Bold',RUST)
        bottom=paragraph(c,leg,27,top,105,11.5,13)
        top=bottom-5
    assert top>=12,top

def rouen_front(c):
    base(c)
    text(c,72,408,'MUSÉE DUCAL',18,'Display',align='center')
    text(c,72,386,'DE ROUEN',18,'Display',align='center');rule(c,372)
    text(c,72,354,'NORMANDIE',9,'Bold',RUST,'center')
    path=NORSE/'TravelMap/Art/Rollo_Chateau_Robert_Le_Diable_Line_v1.png'
    with Image.open(path) as im:iw,ih=im.size
    scale=min(124/iw,190/ih)
    c.drawImage(str(path),(W-iw*scale)/2,153+(190-ih*scale)/2,iw*scale,ih*scale,mask='auto')
    rule(c,135)
    text(c,72,112,'ENTRÉE GÉNÉRALE',12,'Display',align='center')
    text(c,72,94,'GENERAL ADMISSION',9,'Bold',align='center')
    text(c,72,62,'Galerie des mots',12,'Display',RUST,'center')
    text(c,72,44,'Word Gallery',10,align='center')

def rouen_back(c):
    base(c)
    text(c,72,410,'WORD GALLERY',16,'Display',align='center')
    top=paragraph(c,'Match each old word to its new one.',12,392,120,11,13)-7
    # Stack the numbered Saxon list and shuffled Norman list at readable type size.
    for title,items in [('SAXON',SAXON),('NORMAN',NORMAN)]:
        text(c,12,top,title,9,'Bold',RUST);rule(c,top-5)
        top-=20
        for word,tag in items:
            text(c,14,top,str(tag),11,'Bold',RUST)
            text(c,38,top,word,12)
            top-=12.5
        top-=13
    assert top>15,top

def oslo_front(c):
    base(c)
    text(c,72,410,'HNEFATAFL',19,'Display',align='center')
    text(c,72,390,'MUSEUM',17,'Display',align='center')
    text(c,72,374,'OSLO',9,'Bold',RUST,'center');rule(c,364)
    paragraph(c,'the Viking game of kings and corners',16,351,112,9,11)
    c.drawImage(str(HERE/'Hnefatafl/Hnefatafl_Shield_Icon_HandDrawn_v1.png'),55,290,34,34,mask='auto')
    x,y,cell=45,264,18
    c.setStrokeColor(RULE);c.setLineWidth(.5);c.rect(x,y,54,18)
    for j in (1,2):c.line(x+j*cell,y,x+j*cell,y+cell)
    for j in (0,2):c.setFillColor(INK);c.circle(x+(j+.5)*cell,y+9,5,fill=1,stroke=0)
    c.circle(72,y+9,5,stroke=1,fill=0)
    c.setStrokeColor(RUST);c.setLineWidth(.9);c.line(68,y+5,76,y+13);c.line(68,y+13,76,y+5)
    text(c,72,252,'trapped between two: captured',8,align='center')
    paragraphs='\n'.join(RULES).split('\n\n');top=235
    for value in paragraphs:
        top=paragraph(c,value.replace('\n',' '),12,top,120,9.2,10.8)-3
    assert top>25,top
    rule(c,25);text(c,72,13,'GENERAL ADMISSION',8,'Bold',align='center')

def oslo_back(c):base(c)

CONFIG=[
 ('Brattahlid','Museum_Ticket','MuseumTicket',brattahlid_front,brattahlid_back),
 ('Aud','Aud_Ticket','AudTicket',aud_front,aud_back),
 ('Rouen','Rouen_Ticket','RouenTicket',rouen_front,rouen_back),
 ('Oslo','Hnefatafl_Ticket','HnefataflTicket',oslo_front,oslo_back),
]

def build():
    report=[]
    for name,stem,alias,front,back in CONFIG:
        master=PDF/(stem+'_Print.pdf');temporary=WORK/(stem+'_Reflow.pdf')
        c=canvas.Canvas(str(temporary),pagesize=(W,H));c.setTitle(name+' Museum Ticket 2 by 6 inches')
        for fn in (front,back):fn(c);c.showPage()
        c.save();reader=PdfReader(temporary);writer=PdfWriter()
        for page in reader.pages:writer.add_page(page)
        if name=='Oslo':
            # Source grid remains 28.8 points per square, matching the existing map.
            old=PdfReader(SOURCES/'Oslo_Original_Reverse.pdf').pages[0]
            # Make pale paper tones unprinted without changing the original ink
            # image bytes, grid geometry, piece positions or orange corner marks.
            for ref in old['/Resources']['/XObject'].values():
                image=ref.get_object()
                if image.get('/Subtype')=='/Image':
                    assert image['/ColorSpace']=='/DeviceRGB'
                    image[NameObject('/Mask')]=ArrayObject([NumberObject(v) for v in (170,255,145,255,95,255)])
            old.trimbox=RectangleObject([27,20,160,354])
            old.cropbox=RectangleObject([27,20,160,354])
            writer.pages[1].merge_transformed_page(old,Transformation().translate(-21.5,29),expand=False)
        writer.write(master)
        docpdf=pdfium.PdfDocument(str(master))
        sheet=PDF/(stem+'_Letter_Print.pdf')
        c=canvas.Canvas(str(sheet),pagesize=(612,792))
        # Two horizontal copies near the top. Opposite back rotation preserves
        # the original head-to-head portrait faces after long-edge duplexing.
        for i in range(2):
            image=docpdf[i].render(scale=300/72).to_pil()
            image.save(WORK/f'{name}_{i+1}.png')
            image=image.transpose(Image.Transpose.ROTATE_90 if i==0 else Image.Transpose.ROTATE_270)
            for y in (612,450):
                c.drawImage(ImageReader(image),90,y,432,144)
                c.setStrokeColor(INK);c.setLineWidth(.35)
                c.rect(90,y,432,144,fill=0,stroke=1)
            c.showPage()
        c.save()
        d=pdfium.PdfDocument(str(sheet));doc=Document()
        for i in range(2):
            if i:doc.add_section(WD_SECTION.NEW_PAGE)
            sec=doc.sections[-1];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
            image=d[i].render(scale=300/72).to_pil();buf=io.BytesIO();image.save(buf,format='PNG');buf.seek(0)
            positions.add_floating_picture(doc.add_paragraph(),buf,8.5,11,0,0)
        doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
        dest=PRINT/f'Print_Museum_{name}_2x6.docx';doc.save(dest)
        shutil.copy2(dest,PRINT/f'PrintTest_{alias}.docx')
        if name in ('Aud','Brattahlid'):
            shutil.copy2(dest,PRINT/(('Aud' if name=='Aud' else 'Brattahlid')+'_Museum_Ticket_Tan_Paper.docx'))
        assert len(PdfReader(master).pages)==2
        assert all(tuple(p.mediabox)==(0,0,144,432) for p in PdfReader(master).pages)
        report.append({'museum':name,'docx':str(dest),'size_inches':[2,6],'sides':2})
        print(dest)
    (WORK/'build-report.json').write_text(json.dumps(report,indent=2))
    build_pairs()

def build_pairs():
    """One of each museum per paired sheet, with registered duplex backs."""
    stems={name:stem for name,stem,*_ in CONFIG}
    for pair in [('Aud','Rouen'),('Oslo','Brattahlid')]:
        label='_'.join(pair)
        sheet=WORK/f'Museums_{label}_Paired.pdf'
        c=canvas.Canvas(str(sheet),pagesize=(612,792))
        for side in range(2):
            for name,y in zip(pair,(612,450)):
                source=pdfium.PdfDocument(str(PDF/(stems[name]+'_Print.pdf')))
                image=source[side].render(scale=300/72).to_pil()
                image=image.transpose(Image.Transpose.ROTATE_90 if side==0 else Image.Transpose.ROTATE_270)
                c.saveState()
                if side==1:
                    # Hide the source's trim-edge line without scaling the art.
                    # The decorative frame is five points inside and remains.
                    clip=c.beginPath();clip.rect(90+.5,y+.5,432-1,144-1)
                    c.clipPath(clip,stroke=0,fill=0)
                c.drawImage(ImageReader(image),90,y,432,144)
                c.restoreState()
                if side==0:
                    c.setStrokeColor(INK);c.setLineWidth(.35)
                    c.rect(90,y,432,144,fill=0,stroke=1)
            c.showPage()
        c.save()
        pages=pdfium.PdfDocument(str(sheet));doc=Document()
        for side in range(2):
            if side:doc.add_section(WD_SECTION.NEW_PAGE)
            sec=doc.sections[-1];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
            image=pages[side].render(scale=300/72).to_pil()
            buf=io.BytesIO();image.save(buf,format='PNG');buf.seek(0)
            positions.add_floating_picture(doc.add_paragraph(),buf,8.5,11,0,0)
        doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
        dest=PRINT/f'Print_Museums_{label}_2x6.docx'
        doc.save(dest);print(dest)

if __name__=='__main__':build()
