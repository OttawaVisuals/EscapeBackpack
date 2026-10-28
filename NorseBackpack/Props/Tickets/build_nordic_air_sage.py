"""Approved fancy Nordic Air hop tickets on sage stock, with existing calendar dates."""

from pathlib import Path
import ast
import io
import importlib.util
import json
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from ticket_dates import HOP_DATES
from fancy_travel_ticket import draw_ticket

ROOT=Path(__file__).resolve().parents[3]
NORSE=ROOT/'NorseBackpack';HERE=Path(__file__).resolve().parent
OUT=NORSE/'Print';WORK=ROOT.parent/'output/nordic-air-sage'
WORK.mkdir(parents=True,exist_ok=True)
LOGO=HERE/'Nordic_Air_Logo_B_v2.png'
PDF=OUT/'_pdf_source/Nordic_Air_Flights_Sage_Paper.pdf'
DOCX=OUT/'Print_NordicAir_Flights_Sage_Paper.docx'
PREVIEW=HERE/'Nordic_Air_Sage_Final_Preview.png'
W,H=252,144
NAVY,COPPER,SAGE=[HexColor(c) for c in ('#142B36','#98452F','#ADBD98')]
for name,file in [('Sans','SourceSans-400.ttf'),('Semi','SourceSans-600.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(ROOT/'Fonts/CrypticBureau'/file)))
pdfmetrics.registerFont(TTFont('Serif',str(ROOT/'Fonts/CrypticBureau/Cormorant-600.ttf')))


def read_data():
    def literal(file,name):
        tree=ast.parse((HERE/file).read_text(encoding='utf-8'))
        return next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign)
                    and any(isinstance(t,ast.Name) and t.id==name for t in n.targets))
    hops=literal('build_hop_tickets_pdf.py','PROPS')
    result=[]
    for journey,_,_,legs,order in hops:
        for idx in order:
            frm,to,mode,service=legs[idx]
            if mode=='air':result.append(dict(journey=journey,origin=frm,destination=to,service=service,date=HOP_DATES[service]))
    bookings=literal('build_transition_tickets_pdf.py','TICKETS')
    assert [x['service'] for x in result]==['NA 730','NA 722','NA 455','NA 463']
    assert len(bookings)==3
    return result,bookings


def draw_logo(c,x,y,width=43):
    # Clip the transparent padding at placement; do not alter original pixels.
    im=Image.open(LOGO);iw,ih=im.size
    assert im.mode=='RGBA' and im.getpixel((0,0))[3]==0
    bx,by,br,bb=im.getchannel('A').getbbox()
    scale=width/(br-bx);height=(bb-by)*scale
    c.saveState();path=c.beginPath();path.rect(x,y,width,height)
    c.clipPath(path,stroke=0,fill=0)
    c.drawImage(str(LOGO),x-bx*scale,y-(ih-bb)*scale,iw*scale,ih*scale,mask='auto')
    c.restoreState()


def label(c,x,y,value,font='Sans',size=9):
    c.setFillColor(NAVY);c.setFont(font,size);c.drawString(x,y,value)


def ticket(c,rec,x=0,y=0):
    draw_ticket(c,rec,LOGO,'NORDIC','AIR','FLIGHT',COPPER,x,y)


def main():
    records,bookings=read_data()
    c=canvas.Canvas(str(PDF),pagesize=(W,H),pageCompression=1)
    c.setTitle('Nordic Air Flights for Sage Paper');c.setAuthor('The Cryptic Bureau')
    for rec in records:ticket(c,rec);c.showPage()
    c.save()
    proof=WORK/'Sage_Paper_Preview.pdf'
    c=canvas.Canvas(str(proof),pagesize=(W,H))
    c.setFillColor(SAGE);c.rect(0,0,W,H,fill=1,stroke=0)
    ticket(c,records[0]);c.showPage();c.save()
    d=pdfium.PdfDocument(str(proof));d[0].render(scale=4).to_pil().save(PREVIEW)
    reader=PdfReader(PDF)
    assert len(reader.pages)==4
    for page,rec in zip(reader.pages,records):
        value=page.extract_text()
        for expected in [rec['origin'].upper(),rec['destination'].upper(),rec['service'],rec['date'][:6],rec['date'][7:]]:
            assert expected in value
        assert (float(page.mediabox.width),float(page.mediabox.height))==(W,H)
        # Production pages have no full-page paper-colour rectangle.
        assert not any(op == b're' and [float(x) for x in args] == [0,0,W,H]
                       for args,op in page.get_contents().operations)
    sheet=OUT/'_pdf_source/Nordic_Air_Letter_Print.pdf'
    c=canvas.Canvas(str(sheet),pagesize=(612,792),pageCompression=1)
    label(c,54,751,'NORDIC AIR FLIGHT TICKETS','Semi',15)
    label(c,54,733,'Sage green paper. 3½ × 2 inches each. Letter. 100% / actual size.',size=10)
    for i,rec in enumerate(records):
        col,row=i%2,i//2
        ticket(c,rec,43.2+col*(W+21.6),700-H*(row+1)-14*row)
    label(c,43.2,369,'Single-sided. Cut the pale outer rectangles; keep the service/date stub attached.')
    c.setStrokeColor(NAVY);c.setLineWidth(.5);c.line(43.2,340,115.2,340)
    c.line(43.2,336,43.2,344);c.line(115.2,336,115.2,344)
    label(c,128,337,'This line must measure exactly 1 inch before cutting.')
    c.showPage();c.save()
    # Render every production ticket and the printable sheet for visual verification.
    for i,page in enumerate(pdfium.PdfDocument(str(PDF))):
        page.render(scale=4).to_pil().save(WORK/f'ticket-{i+1}.png')
    pdfium.PdfDocument(str(sheet))[0].render(scale=2).to_pil().save(WORK/'letter-preview.png')
    pdf=pdfium.PdfDocument(str(sheet));image=pdf[0].render(scale=300/72).to_pil()
    stream=io.BytesIO();image.save(stream,format='PNG');stream.seek(0)
    spec=importlib.util.spec_from_file_location('positions',NORSE/'Postcards/build_print_test_docx.py')
    pos=importlib.util.module_from_spec(spec);spec.loader.exec_module(pos)
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    pos.add_floating_picture(doc.add_paragraph(),stream,8.5,11,0,0)
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title='Nordic Air Flights for Sage Paper';doc.core_properties.author='The Cryptic Bureau'
    doc.save(DOCX)
    (WORK/'verified_data.json').write_text(json.dumps({'hop_flights':records,'connecting_bookings_retained':bookings,
        'source':'current builders, confirmed by designer table','background_fill':'none in production',
        'size_inches':[3.5,2],'physical_test':'not performed'},indent=2))
    print(DOCX);print(PREVIEW)


if __name__=='__main__':main()
