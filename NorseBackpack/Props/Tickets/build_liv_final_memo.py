"""Current final-letter wording on approved 4 x 6 Bureau Archive stationery."""
from pathlib import Path
import ast
import importlib.util
import io

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, ArrayObject, NumberObject
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[3]
NORSE = ROOT / 'NorseBackpack'
WORK = ROOT.parent / 'output/liv-final-memo'
WORK.mkdir(parents=True, exist_ok=True)
PDF = NORSE / 'Print/_pdf_source/Liv_Final_Memo_4x6_Print.pdf'
DOCX = NORSE / 'Print/Print_LivFinalMemo_4x6.docx'
LOGO = ROOT / 'Brand/CrypticBureau/LogoConcepts/CB_Logo_B_Archive_Hidden_Flaps_v1.png'
RAVEN = NORSE / 'Art/Raven/Raven_Profile_Knotwork_Pencil_v2.png'
tree = ast.parse(Path(__file__).with_name('build_liv_final_letter_pdf.py').read_text(encoding='utf-8'))
PARAS = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
             and any(isinstance(t, ast.Name) and t.id == 'PARAS' for t in n.targets))
PARAS += ['With love,', 'Aunt Liv']
pdfmetrics.registerFont(TTFont('LivHand', str(ROOT/'Fonts/Nothing_You_Could_Do/NothingYouCouldDo-Regular.ttf')))
pdfmetrics.registerFont(TTFont('Bureau', str(ROOT/'Fonts/CrypticBureau/Cormorant-600.ttf')))
BURGUNDY, COPPER, INK = [HexColor(c) for c in ('#542C3B','#B87650','#283B34')]


def emblem(c, x, y, width):
    # Clip the original approved envelope, retaining its geometry and hidden
    # flap lettering; use a PDF colour key for its ivory background/seams.
    sx, sy, sw, sh = 280, 470, 700, 450
    scale = width / sw
    height = sh * scale
    c.saveState()
    path = c.beginPath(); path.rect(x, y, width, height)
    c.clipPath(path, stroke=0, fill=0)
    c.drawImage(str(LOGO), x-sx*scale, y-sy*scale, 1254*scale, 1254*scale,
                mask=[225,255,210,255,190,255])
    c.restoreState()


def main():
    raw = WORK/'memo_layout.pdf'
    c = canvas.Canvas(str(raw), pagesize=(288,432), pageCompression=1)
    c.setTitle('Liv Final Memo 4 by 6'); c.setAuthor('The Cryptic Bureau')
    c.setStrokeColor(COPPER); c.setLineWidth(.5)
    c.line(16,46,16,413)
    c.line(25,373,58,373)
    emblem(c,25,387,39)
    c.setFillColor(BURGUNDY)
    c.setFont('Bureau',9); c.drawString(72,408,'THE')
    c.setFont('Bureau',13); c.drawString(72,394,'CRYPTIC BUREAU')
    c.drawImage(str(RAVEN),216,365,57,57,mask='auto',preserveAspectRatio=True)
    emblem(c,16,16,28)
    size, leading, x, width, y = 12.5, 16, 27, 240, 346
    c.setFillColor(INK); c.setFont('LivHand',size)
    lines = []
    for i, paragraph in enumerate(PARAS):
        words = paragraph.split(); row = ''
        wrapped = []
        for word in words:
            trial = (row+' '+word).strip()
            if row and pdfmetrics.stringWidth(trial,'LivHand',size) > width:
                wrapped.append(row); row=word
            else: row=trial
        wrapped.append(row)
        for row in wrapped:
            c.drawString(x,y,row); lines.append(row); y-=leading
        if i not in (0,len(PARAS)-2): y-=7
    assert y > 48, y
    c.showPage(); c.save()
    # Black stamp impression from original alpha art, without regenerating it.
    reader=PdfReader(raw); writer=PdfWriter(); page=writer.add_page(reader.pages[0])
    for ref in page['/Resources']['/XObject'].values():
        obj=ref.get_object()
        if '/SMask' in obj:
            obj[NameObject('/Decode')]=ArrayObject([NumberObject(0)]*6)
    writer.write(PDF)
    check=PdfReader(PDF).pages[0].extract_text()
    for paragraph in PARAS:
        assert ' '.join(paragraph.split()) in ' '.join(check.split())
    pdf=pdfium.PdfDocument(str(PDF))
    image=pdf[0].render(scale=300/72).to_pil()
    image.save(WORK/'memo_preview.png')
    stream=io.BytesIO(); image.save(stream,format='PNG'); stream.seek(0)
    spec=importlib.util.spec_from_file_location('positions',NORSE/'Postcards/build_print_test_docx.py')
    pos=importlib.util.module_from_spec(spec); spec.loader.exec_module(pos)
    doc=Document(); sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    p=doc.add_paragraph(); pos.add_floating_picture(p,stream,4,6,.6,.3)
    shape=p._p.xpath('.//pic:spPr')[0]
    line=OxmlElement('a:ln');line.set('w','3175')
    solid=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','000000')
    solid.append(color);line.append(solid);shape.append(line)
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title='Liv Final Memo 4 by 6'
    doc.core_properties.author='The Cryptic Bureau'
    doc.save(DOCX)
    print('Text unchanged; final text bottom:',y+leading)
    print(DOCX)


if __name__=='__main__': main()
