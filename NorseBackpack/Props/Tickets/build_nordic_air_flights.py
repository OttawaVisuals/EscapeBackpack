"""Four flight hop tickets in the approved modern Nordic Air direction.

Retains exact source endpoints and their scrambled print order. Existing mixed
hop-ticket masters remain unchanged while the other transport types are designed.
"""
from pathlib import Path
import ast
import io
import importlib.util
import json

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import black, HexColor
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement

ROOT=Path(__file__).resolve().parents[3]
NORSE=ROOT/'NorseBackpack'
HERE=Path(__file__).resolve().parent
WORK=ROOT.parent/'output/nordic-air-final'
WORK.mkdir(parents=True,exist_ok=True)
OUT=NORSE/'Print'
PDF=OUT/'_pdf_source/Nordic_Air_Flights_Redesign.pdf'
DOCX=OUT/'Print_NordicAir_Flights.docx'
PREVIEW=HERE/'Nordic_Air_Final_Preview.png'
W,H=288,126
for font,file in [('Sans','SourceSans-400.ttf'),('Semi','SourceSans-600.ttf')]:
    pdfmetrics.registerFont(TTFont(font,str(ROOT/'Fonts/CrypticBureau'/file)))


def flights():
    tree=ast.parse((HERE/'build_hop_tickets_pdf.py').read_text(encoding='utf-8'))
    props=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign)
               and any(isinstance(t,ast.Name) and t.id=='PROPS' for t in n.targets))
    result=[]
    for key,_,_,hops,order in props:
        for idx in order:
            frm,to,mode,_=hops[idx]
            if mode=='air':result.append((key,frm,to))
    assert len(result)==4
    return result


def plane(c,x,y):
    # Original geometric airplane silhouette, no external icon/font dependency.
    points=[(0,7),(7,8),(13,16),(16,16),(12,8),(19,8),(22,11),(24,11),
            (22,6),(24,1),(22,1),(19,4),(12,4),(16,-4),(13,-4),(7,4),(0,5)]
    c.saveState();c.translate(x,y);c.scale(.68,.68)
    p=c.beginPath();p.moveTo(*points[0])
    for pt in points[1:]:p.lineTo(*pt)
    p.close();c.setFillColor(black);c.drawPath(p,fill=1,stroke=0);c.restoreState()


def ticket(c,frm,to,x=0,y=0):
    c.saveState();c.translate(x,y)
    c.setStrokeColor(HexColor('#888888'));c.setLineWidth(.4)
    c.rect(.2,.2,W-.4,H-.4,fill=0,stroke=1)
    plane(c,12,105)
    c.setFillColor(black);c.setFont('Semi',12)
    c.drawString(36,107,'NORDIC AIR')
    c.setFont('Sans',9);c.drawRightString(266,108,'FLIGHT')
    c.setStrokeColor(HexColor('#999999'));c.setLineWidth(.4)
    c.line(12,97,269,97);c.line(12,22,269,22)
    c.setDash(2,2);c.line(276,3,276,H-3);c.setDash([])
    for label,value,ly,vy in [('FROM',frm,84,66),('TO',to,49,31)]:
        c.setFillColor(black);c.setFont('Sans',9);c.drawString(12,ly,label)
        c.setFont('Semi',14)
        assert pdfmetrics.stringWidth(value.upper(),'Semi',14)<252
        c.drawString(12,vy,value.upper())
    c.setFont('Sans',9);c.drawString(12,9,'PASSENGER')
    c.setFont('Semi',9);c.drawString(69,9,'LIV ERICSON')
    c.restoreState()


def main():
    data=flights()
    c=canvas.Canvas(str(PDF),pagesize=(W,H),pageCompression=1)
    c.setTitle('Nordic Air Flight Hop Tickets');c.setAuthor('The Cryptic Bureau')
    for _,frm,to in data:ticket(c,frm,to);c.showPage()
    c.save()
    reader=PdfReader(PDF)
    assert len(reader.pages)==4
    for page,(_,frm,to) in zip(reader.pages,data):
        value=page.extract_text()
        assert frm.upper() in value and to.upper() in value
        assert float(page.mediabox.width)==W and float(page.mediabox.height)==H
    d=pdfium.PdfDocument(str(PDF));d[0].render(scale=4).to_pil().save(PREVIEW)
    sheet=WORK/'Nordic_Air_Letter_Print.pdf'
    c=canvas.Canvas(str(sheet),pagesize=(612,792),pageCompression=1)
    c.setTitle('Nordic Air Flight Tickets Letter Print')
    c.setFillColor(black);c.setFont('Semi',15);c.drawString(54,751,'NORDIC AIR FLIGHT TICKETS')
    c.setFont('Sans',10);c.drawString(54,733,'4 × 1¾ inches each. Letter paper. 100% / actual size. Single-sided.')
    for i,(_,frm,to) in enumerate(data):ticket(c,frm,to,162,700-H*(i+1)-14*i)
    c.setFont('Sans',9);c.drawString(54,117,'Cut along each outer border. The dashed stub divider is decorative; do not cut it.')
    c.setStrokeColor(black);c.setLineWidth(.5);c.line(54,88,126,88)
    c.line(54,84,54,92);c.line(126,84,126,92)
    c.drawString(138,85,'Check this line measures 1 inch before cutting.')
    c.showPage();c.save()
    pdf=pdfium.PdfDocument(str(sheet))
    im=pdf[0].render(scale=300/72).to_pil()
    stream=io.BytesIO();im.save(stream,format='PNG');stream.seek(0)
    spec=importlib.util.spec_from_file_location('positions',NORSE/'Postcards/build_print_test_docx.py')
    pos=importlib.util.module_from_spec(spec);spec.loader.exec_module(pos)
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    p=doc.add_paragraph();pos.add_floating_picture(p,stream,8.5,11,0,0)
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title='Nordic Air Flight Tickets'
    doc.core_properties.author='The Cryptic Bureau'
    doc.save(DOCX)
    (WORK/'checks.json').write_text(json.dumps({'tickets':data,'size_inches':[4,1.75],
        'fonts_points':{'carrier':12,'places':14,'labels_and_passenger':9},
        'source_endpoints':'unchanged','source_scramble_order':'preserved','physical_test':'not performed'},indent=2))
    print(DOCX);print(PREVIEW)


if __name__=='__main__':main()
