"""Four fancy Longship Ferries hop tickets for dark-green paper. Each ticket carries its departure date (ticket_dates.py)."""

from pathlib import Path
import ast
import io
import importlib.util
import json
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from ticket_dates import HOP_DATES
from fancy_travel_ticket import draw_ticket

ROOT=Path(__file__).resolve().parents[3]
NORSE=ROOT/'NorseBackpack';HERE=Path(__file__).resolve().parent
OUT=NORSE/'Print';WORK=ROOT.parent/'output/longship-ferries'
WORK.mkdir(parents=True,exist_ok=True)
PDF=OUT/'_pdf_source/Longship_Ferries_Green_Paper.pdf'
DOCX=OUT/'Print_LongshipFerries_Green_Paper.docx'
PREVIEW=HERE/'Longship_Ferries_Green_Final_Preview.png'
LOGO=HERE/'Longship_Ferries_Logo_A_v3.png'
W,H=252,144
INK,ACCENT,GREEN=[HexColor(c) for c in ('#142B36','#142B36','#678468')]
for name,file in [('Sans','SourceSans-400.ttf'),('Semi','SourceSans-600.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(ROOT/'Fonts/CrypticBureau'/file)))


def read_data():
    tree=ast.parse((HERE/'build_hop_tickets_pdf.py').read_text(encoding='utf-8'))
    props=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign)
               and any(isinstance(t,ast.Name) and t.id=='PROPS' for t in n.targets))
    result=[]
    for journey,_,_,legs,order in props:
        for idx in order:
            frm,to,mode,service=legs[idx]
            if mode=='ferry':result.append(dict(journey=journey,origin=frm,destination=to,service=service,date=HOP_DATES[service]))
    assert [x['service'] for x in result]==['LF 11','LF 204','LF 810','LF 377']
    return result


def draw_logo(c,x,y,width=29):
    im=Image.open(LOGO)
    assert im.mode=='RGBA' and im.getpixel((0,0))[3]==0
    # Fit original transparent bitmap into the header, without modifying it.
    bx,by,br,bb=im.getchannel('A').point(lambda v:255 if v>128 else 0).getbbox()
    scale=min(width/(br-bx),24/(bb-by))
    iw,ih=im.size
    ox=x+(width-(br-bx)*scale)/2
    oy=y+(24-(bb-by)*scale)/2
    c.saveState()
    p=c.beginPath();p.rect(x,y,width,24);c.clipPath(p,stroke=0,fill=0)
    c.drawImage(str(LOGO),ox-bx*scale,oy-(ih-bb)*scale,iw*scale,ih*scale,mask='auto')
    c.restoreState()


def text(c,x,y,value,font='Sans',size=9,color=INK):
    c.setFillColor(color);c.setFont(font,size);c.drawString(x,y,value)


def ticket(c,rec,x=0,y=0):
    draw_ticket(c,rec,LOGO,'LONGSHIP','FERRIES','FERRY',HexColor('#A74F34'),x,y)


def main():
    records=read_data()
    c=canvas.Canvas(str(PDF),pagesize=(W,H),pageCompression=1)
    c.setTitle('Longship Ferries Tickets for Dark Green Paper');c.setAuthor('The Cryptic Bureau')
    for rec in records:ticket(c,rec);c.showPage()
    c.save()
    proof=WORK/'Green_Paper_Preview.pdf'
    c=canvas.Canvas(str(proof),pagesize=(W,H))
    c.setFillColor(GREEN);c.rect(0,0,W,H,fill=1,stroke=0)
    ticket(c,records[0]);c.showPage();c.save()
    d=pdfium.PdfDocument(str(proof));d[0].render(scale=4).to_pil().save(PREVIEW)
    reader=PdfReader(PDF)
    assert len(reader.pages)==4
    for page,rec in zip(reader.pages,records):
        value=page.extract_text()
        for expected in [rec['origin'].upper(),rec['destination'].upper(),rec['service'],rec['date'][:6],rec['date'][7:]]:
            assert expected in value
        assert (float(page.mediabox.width),float(page.mediabox.height))==(W,H)
        assert not any(op==b're' and [float(x) for x in args]==[0,0,W,H]
                       for args,op in page.get_contents().operations)
    sheet=OUT/'_pdf_source/Longship_Ferries_Letter_Print.pdf'
    c=canvas.Canvas(str(sheet),pagesize=(612,792),pageCompression=1)
    text(c,43.2,751,'LONGSHIP FERRIES TICKETS','Semi',15)
    text(c,43.2,733,'Dark-green paper. 3½ × 2 inches each. Letter. 100% / actual size.',size=10)
    for i,rec in enumerate(records):
        col,row=i%2,i//2
        ticket(c,rec,43.2+col*(W+21.6),700-H*(row+1)-14*row)
    text(c,43.2,369,'Single-sided. Cut the pale outer rectangles; keep each service/date stub attached.')
    c.setStrokeColor(INK);c.setLineWidth(.5);c.line(43.2,340,115.2,340)
    c.line(43.2,336,43.2,344);c.line(115.2,336,115.2,344)
    text(c,128,337,'This line must measure exactly 1 inch before cutting.')
    c.showPage();c.save()
    pdfium.PdfDocument(str(sheet))[0].render(scale=2).to_pil().save(WORK/'letter-preview.png')
    pdf=pdfium.PdfDocument(str(sheet));im=pdf[0].render(scale=300/72).to_pil()
    stream=io.BytesIO();im.save(stream,format='PNG');stream.seek(0)
    spec=importlib.util.spec_from_file_location('positions',NORSE/'Postcards/build_print_test_docx.py')
    pos=importlib.util.module_from_spec(spec);spec.loader.exec_module(pos)
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    pos.add_floating_picture(doc.add_paragraph(),stream,8.5,11,0,0)
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title='Longship Ferries Tickets for Dark Green Paper'
    doc.core_properties.author='The Cryptic Bureau';doc.save(DOCX)
    (WORK/'verified_data.json').write_text(json.dumps({'tickets':records,'size_inches':[3.5,2],
        'paper_hex':'#678468','font_points':{'places_range':[14,18],'wordmark':12.5,'labels':8,'service_range':[10,13],'date':11.5},
        'background':'none in production','physical_test':'not performed'},indent=2))
    print(DOCX);print(PREVIEW)


if __name__=='__main__':main()

