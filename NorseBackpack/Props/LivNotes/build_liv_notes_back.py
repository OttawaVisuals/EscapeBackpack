"""Attach the approved 20 Oct help reverse to Liv's existing 4 x 6 inch front.

The adopted front is preserved in Sources/Liv_Notes_Front_Master.pdf.
Run with --refresh-front only after deliberately replacing the live front PDF.
Produces the two-page card PDF and aligned two-page Letter Word print sheet.
"""
from pathlib import Path
import argparse
import importlib.util
import io
import json
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from docx.oxml import OxmlElement
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, ByteStringObject, TextStringObject
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[1]
PRINT = PROJECT / 'Print'
PDF = PRINT / '_pdf_source/Liv_Notes_Card_Print.pdf'
WORD = PRINT / 'Print_LivNotesCard.docx'
FRONT = HERE / 'Sources/Liv_Notes_Front_Master.pdf'
BACK = HERE / 'Liv_Notes_Back.png'
ARCHIVE = PROJECT.parents[1] / 'Norse_Archive/NorseBackpack/Props/LivNotes/Before_Help_Back_20261020'


def archive_once():
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    for src in (PDF, WORD):
        dest = ARCHIVE / src.name
        if src.exists() and not dest.exists():
            shutil.copy2(src, dest)
    readme = ARCHIVE / 'README.md'
    if not readme.exists():
        readme.write_text('20 Oct 2026: preserved the single-sided Liv notes PDF and Word master before adopting the approved help and code-guide reverse. Front unchanged. Current outputs keep their original filenames.\n', encoding='utf-8')


def build(refresh_front=False):
    archive_once()
    if refresh_front or not FRONT.exists():
        first = PdfReader(PDF)
        w = PdfWriter()
        w.add_page(first.pages[0])
        w.write(FRONT)
    front = PdfReader(FRONT)
    original_text = front.pages[0].extract_text()
    # Remove only the full-card cream fill. Retain the text, ink and fine border.
    stream = front.pages[0].get_contents()
    operations = stream.operations
    removed = 0
    for n, (args, op) in enumerate(operations):
        if op == b're' and tuple(float(x) for x in args) == (0, 0, 288, 432):
            if operations[n+1][1] in (b'f', b'f*'):
                operations[n+1] = ([], b'n')
                removed += 1
    # Preserve subset-font character codes during stream serialization.
    def raw_bytes(value):
        if isinstance(value, TextStringObject):
            return ByteStringObject(value.original_bytes)
        if isinstance(value, ArrayObject):
            return ArrayObject([raw_bytes(x) for x in value])
        return value
    stream.operations = [([raw_bytes(x) for x in args], op) for args, op in operations]
    if removed:
        clean = PdfWriter()
        clean_page = clean.add_page(front.pages[0])
        clean_page.replace_contents(stream)
        clean.write(FRONT)
        front = PdfReader(FRONT)
    assert front.pages[0].extract_text() == original_text
    width, height = map(float, (front.pages[0].mediabox.width, front.pages[0].mediabox.height))
    assert (width, height) == (288.0, 432.0)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    c.setTitle('Liv notes card help and code guide reverse')
    c.drawImage(ImageReader(str(BACK)), 0, 0, width=width, height=height, mask='auto')
    c.save()
    out = PdfWriter()
    out.add_page(front.pages[0])
    out.add_page(PdfReader(buf).pages[0])
    out.add_metadata({'/Title': 'Liv notes card front and help reverse'})
    out.write(PDF)

    # Keep both 4 x 6 inch faces aligned in the printer's upper page area.
    spec = importlib.util.spec_from_file_location('positions', PROJECT / 'Postcards/build_print_test_docx.py')
    positions = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(positions)
    doc = Document()
    rendered = pdfium.PdfDocument(str(PDF))
    for index in range(len(rendered)):
        if index:
            doc.add_section(WD_SECTION.NEW_PAGE)
        section = doc.sections[-1]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
            setattr(section, side, Inches(.25))
        img = rendered[index].render(scale=300/72).to_pil()
        img.save(HERE / ('Liv_Notes_Card_Preview.png' if index == 0 else 'Liv_Notes_Back_Preview.png'))
        stream = io.BytesIO()
        img.save(stream, format='PNG')
        stream.seek(0)
        paragraph = doc.add_paragraph()
        positions.add_floating_picture(paragraph, stream, 4, 6, 2.25, .5)
        if index == 0:
            # Front-only cut guide at the actual card edge, outside its decoration.
            props = paragraph._p.xpath('.//pic:spPr')[0]
            line = OxmlElement('a:ln'); line.set('w', '6350')  # 0.5 pt
            fill = OxmlElement('a:solidFill')
            colour = OxmlElement('a:srgbClr'); colour.set('val', '777777')
            fill.append(colour); line.append(fill)
            dash = OxmlElement('a:prstDash'); dash.set('val', 'dash')
            line.append(dash); props.append(line)
    doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    doc.core_properties.title = 'Liv Notes Card Front and Help Reverse'
    doc.core_properties.comments = 'Two Letter pages. Print at Actual size, duplex, flip on long edge. Both 4 x 6 inch faces at x 2.25 inches, y 0.5 inches. No printed background on either face. Test alignment and QR scanning before batch printing.'
    doc.save(WORD)
    checks = {'pdf': str(PDF), 'word': str(WORD), 'card_inches': [4,6], 'word_position_inches': [2.25,.5], 'front_background_removed': True, 'pages': len(PdfReader(PDF).pages), 'front_text_preserved': front.pages[0].extract_text() == PdfReader(PDF).pages[0].extract_text(), 'print': 'Actual size, duplex, long edge; physical scan and alignment untested'}
    (HERE / 'Sources/Liv_Notes_Back_Checks.json').write_text(json.dumps(checks, indent=2), encoding='utf-8')
    print(json.dumps(checks, indent=2))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--refresh-front', action='store_true')
    build(ap.parse_args().refresh_front)
