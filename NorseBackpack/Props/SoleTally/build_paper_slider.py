"""Letter print templates and illustrated Word guide for the paper tally prototype.

No AI artwork: dimensions, digits and assembly diagrams are drawn deterministically.
The original 3D-print tally is retained. Physical fit still requires a paper test.
"""
from pathlib import Path
import io
import json
import sys

from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, black, white
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
NORSE = ROOT/'NorseBackpack'
OUT = NORSE/'Print'
COMPACT = '--compact' in sys.argv
SUFFIX = '_Compact_4x6_Envelope' if COMPACT else ''
WORK = ROOT.parent/('output/paper-tally-compact' if COMPACT else 'output/paper-tally')
WORK.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)
I = 72
W,H = (4 if COMPACT else 6)*I,(2 if COMPACT else 2.5)*I
STRIP_W,STRIP_L = (.55 if COMPACT else .75)*I,(6 if COMPACT else 8)*I
CAP_W,CAP_H = (.85 if COMPACT else 1.15)*I,(.125 if COMPACT else .2)*I
PITCH,FIRST = ((.4,1.2) if COMPACT else (.55,1.525))
PITCH,FIRST = PITCH*I,FIRST*I
DIGIT_H = (.3 if COMPACT else .4)*I
RAIL_PITCH = (.9 if COMPACT else 1.35)*I
RAIL_W = (.25 if COMPACT else .45)*I
RAILS = [.075*I+RAIL_PITCH*k for k in range(5)]
CENTRES = [(RAILS[k]+RAIL_W+RAILS[k+1])/2 for k in range(4)]
WINDOW_W,WINDOW_H = (.38 if COMPACT else .58)*I,(.38 if COMPACT else .5)*I
SEGMENTS = dict(zip('0123456789', ['abcdef','bc','abdeg','abcdg','bcfg','acdfg','acdefg','abc','abcdefg','abcdfg']))
SEGMENTS.update({'S':'acdfg','O':'abcdef','L':'def','E':'adefg'})
ROTATE = dict(zip('abcdefg','defabcg'))
TAN,SAGE,GREY = [HexColor(s) for s in ('#E5CCAA','#A9BD9A','#666666')]
pdfmetrics.registerFont(TTFont('Cinzel',str(ROOT/'Fonts/Cinzel/static/Cinzel-Bold.ttf')))


def text(c,x,y,s,size=10,font='Helvetica',color=black):
    c.setFillColor(color);c.setFont(font,size);c.drawString(x,y,s)


def digit(c,x,y,char,height=DIGIT_H):
    w,h,t = .6*height,height,.13*height
    pos = {'a':(-w/2,h/2-t,w,t),'d':(-w/2,-h/2,w,t),'g':(-w/2,-t/2,w,t),
           'b':(w/2-t,0,t,h/2),'c':(w/2-t,-h/2,t,h/2),
           'e':(-w/2,-h/2,t,h/2),'f':(-w/2,0,t,h/2)}
    c.setFillColor(black)
    for s in SEGMENTS[char]:
        xx,yy,ww,hh=pos[s]; c.rect(x+xx,y+yy,ww,hh,fill=1,stroke=0)


def strip_path(c,x,y):
    # y is the bottom of the complete strip, including its two end tabs.
    inset=(CAP_W-STRIP_W)/2
    p=c.beginPath();p.moveTo(x,y);p.lineTo(x+CAP_W,y)
    p.lineTo(x+CAP_W,y+CAP_H);p.lineTo(x+inset+STRIP_W,y+CAP_H)
    p.lineTo(x+inset+STRIP_W,y+STRIP_L-CAP_H)
    p.lineTo(x+CAP_W,y+STRIP_L-CAP_H);p.lineTo(x+CAP_W,y+STRIP_L)
    p.lineTo(x,y+STRIP_L);p.lineTo(x,y+STRIP_L-CAP_H)
    p.lineTo(x+inset,y+STRIP_L-CAP_H);p.lineTo(x+inset,y+CAP_H)
    p.lineTo(x,y+CAP_H);p.close()
    return p


def face(c,x,y,schematic=False):
    c.setStrokeColor(black);c.setLineWidth(.6)
    if schematic:c.setFillColor(TAN)
    c.rect(x,y,W,H,stroke=1,fill=int(schematic))
    c.rect(x+8,y+8,W-16,H-16,stroke=1,fill=0)
    c.setFillColor(black);c.setFont('Cinzel',17)
    c.drawCentredString(x+W/2,y+H-33,'TREASURE TALLY')
    for xx in CENTRES:
        c.setDash([] if schematic else [3,2])
        c.rect(x+xx-WINDOW_W/2,y+H/2-WINDOW_H/2,WINDOW_W,WINDOW_H,stroke=1,fill=0)
    c.setDash([]);c.setFont('Helvetica',11)
    c.drawCentredString(x+W/2,y+37,'Set your total.')
    # Small angular corner details without clue-bearing runes or numbers.
    for dx,dy,sx,sy in [(13,13,1,1),(W-13,13,-1,1),(13,H-13,1,-1),(W-13,H-13,-1,-1)]:
        p=c.beginPath();p.moveTo(x+dx,y+dy+sy*9);p.lineTo(x+dx,y+dy)
        p.lineTo(x+dx+sx*9,y+dy);c.drawPath(p)


def backing(c,x,y,schematic=False):
    c.setStrokeColor(black);c.setLineWidth(.6)
    if schematic:c.setFillColor(TAN)
    c.rect(x,y,W,H,stroke=1,fill=int(schematic))
    for xx in RAILS:
        c.setDash([2,2]);c.setStrokeColor(GREY)
        c.rect(x+xx,y,RAIL_W,H,fill=0,stroke=1)
        c.setDash([])
        if schematic:
            c.setFillColor(HexColor('#C7B79D'));c.rect(x+xx,y,RAIL_W,H,fill=1,stroke=0)
        else:
            c.saveState();c.translate(x+xx+RAIL_W/2,y+H/2);c.rotate(90)
            c.setFont('Helvetica',7);c.setFillColor(GREY);c.drawCentredString(0,-2,'GLUE RAIL STACK HERE');c.restoreState()
    c.setDash([]);c.setStrokeColor(black)


def scale_check(c,y):
    c.setLineWidth(.7);c.setStrokeColor(black)
    c.line(54,y,126,y);c.line(54,y-4,54,y+4);c.line(126,y-4,126,y+4)
    text(c,139,y-3,'This line must measure exactly 1 inch (25.4 mm).',9)


def templates():
    path=WORK/'Paper_Tally_Templates.pdf'
    c=canvas.Canvas(str(path),pagesize=(612,792),pageCompression=1)
    c.setTitle('Paper treasure tally templates')
    text(c,54,750,('COMPACT PAPER TALLY' if COMPACT else 'PAPER TREASURE TALLY')+'  |  1 OF 2',16,'Helvetica-Bold')
    text(c,54,732,'Print on tan cardstock. Letter paper. 100% / actual size. Single-sided.',10)
    text(c,90,716,'FRONT - cut outer edge; cut out the four dashed windows.',9,'Helvetica-Bold')
    face(c,90,527)
    text(c,90,506,'BACKING - cut outer edge only. Keep the five dashed rail zones.',9,'Helvetica-Bold')
    backing(c,90,317)
    text(c,54,289,'TEN RAIL PIECES - cut all ten; glue in pairs to make five stacks.',9,'Helvetica-Bold')
    for i in range(10):
        x=69+i*48
        c.setLineWidth(.6);c.rect(x,96,RAIL_W,H,fill=0,stroke=1)
        text(c,x+9,105,str(i+1),7,color=GREY)
    text(c,54,71,'Solid outer lines = cut. Dashed windows = remove. Backing guides = do not cut.',9)
    scale_check(c,48)
    c.showPage()
    text(c,54,750,('COMPACT PAPER TALLY' if COMPACT else 'PAPER TREASURE TALLY')+'  |  2 OF 2',16,'Helvetica-Bold')
    text(c,54,732,'Print on light green or white cardstock. 100% / actual size. Single-sided.',10)
    text(c,54,716,'Cut four complete strips. Keep BOTH wider end tabs. No folding or joining.',9,'Helvetica-Bold')
    bottom=96
    for i in range(4):
        x=(.7+2*i)*I
        c.setStrokeColor(black);c.setLineWidth(.6)
        c.drawPath(strip_path(c,x,bottom),stroke=1,fill=0)
        for n in range(10):digit(c,x+CAP_W/2,bottom+STRIP_L-FIRST-n*PITCH,str(n))
        text(c,x+10,bottom-18,f'SLIDER {i+1}',8,'Helvetica-Bold')
    scale_check(c,48)
    c.showPage();c.save()
    return path


def fig(name,draw,width=600,height=210):
    path=WORK/(name+'.pdf');c=canvas.Canvas(str(path),pagesize=(width,height))
    draw(c,width,height);c.showPage();c.save()
    d=pdfium.PdfDocument(str(path));png=WORK/(name+'.png')
    d[0].render(scale=2).to_pil().save(png)
    return png


def figures():
    def parts(c,w,h):
        c.saveState();c.translate(14,105);c.scale(.58,.58);face(c,0,0,True);c.restoreState()
        text(c,14,96,'Front with four windows',11,'Helvetica-Bold')
        c.saveState();c.translate(14,15);c.scale(.58,.35);backing(c,0,0,True);c.restoreState()
        text(c,14,3,'Backing with rail guides',10)
        for i in range(4):
            x=335+i*55;c.setFillColor(SAGE);c.rect(x,20,28,180,stroke=1,fill=1)
            for n in range(10):digit(c,x+14,182-n*16,str(n),11)
        text(c,330,3,'4 full number strips',10)
    def cut(c,w,h):
        c.saveState();c.translate(20,33);c.scale(.85,.85);face(c,0,0,True);c.restoreState()
        for xx in CENTRES:
            text(c,20+xx*.85-10,33+H*.85/2-3,'CUT',7,'Helvetica-Bold')
        text(c,421,153,'Remove only',12,'Helvetica-Bold');text(c,421,136,'the four windows.',11)
        text(c,421,107,'Use a ruler and knife',10);text(c,421,92,'on a cutting mat.',10)
        text(c,421,57,'Keep the backing solid.',10,'Helvetica-Bold')
    def rails(c,w,h):
        c.saveState();c.translate(15,33);c.scale(.85,.85);backing(c,0,0,True);c.restoreState()
        text(c,417,161,'5 rail stacks',12,'Helvetica-Bold')
        text(c,417,143,'2 pieces per stack',11)
        text(c,417,111,'Glue on the dashed',10);text(c,417,96,'zones only.',10)
        text(c,417,62,'4 clear channels',11,'Helvetica-Bold')
        for xx in CENTRES:
            c.setStrokeColor(GREY);c.line(15+xx*.85,29,15+xx*.85,13)
    def layers(c,w,h):
        # Exaggerated layer heights for an explanatory side section.
        c.setFillColor(TAN);c.rect(35,30,510,20,fill=1,stroke=1)
        c.setFillColor(HexColor('#C7B79D'));c.rect(35,50,70,27,fill=1,stroke=1);c.rect(475,50,70,27,fill=1,stroke=1)
        c.setFillColor(SAGE);c.rect(140,51,300,10,fill=1,stroke=1)
        c.setFillColor(TAN);c.rect(35,77,510,16,fill=1,stroke=1)
        text(c,35,111,'Front',10,'Helvetica-Bold');text(c,235,111,'Slider moves in the gap',10,'Helvetica-Bold')
        text(c,35,14,'Backing',9);text(c,470,14,'Rail stack',9)
        text(c,35,132,'Side view - layer thickness exaggerated',11,'Helvetica-Bold')
    def install(c,w,h):
        # True relative strip/sleeve length at 0; clipped only by diagram bounds.
        sc=.29;ox,oy=40,25
        for xx in CENTRES:
            x=ox+xx*sc-CAP_W*sc/2
            c.saveState();c.translate(x,oy);c.scale(sc,sc)
            c.setFillColor(SAGE);c.drawPath(strip_path(c,0,0),fill=1,stroke=1)
            for n in range(10):digit(c,CAP_W/2,STRIP_L-FIRST-n*PITCH,str(n))
            c.restoreState()
        by=oy+(STRIP_L-FIRST-H/2)*sc
        c.saveState();c.translate(ox,by);c.scale(sc,sc)
        # Transparent backing outline with rails; strips remain visible in diagram.
        c.setStrokeColor(black);c.rect(0,0,W,H,fill=0,stroke=1)
        for x in RAILS:
            c.setFillColor(TAN);c.rect(x,0,RAIL_W,H,stroke=1,fill=1)
        c.restoreState()
        text(c,260,174,'Lay one strip in each channel.',12,'Helvetica-Bold')
        text(c,260,149,'All digits face up; 0 is nearest the top.',11)
        text(c,260,126,'Leave both wider tabs outside the case.',11)
        text(c,260,94,'Do not glue or tape the sliders.',11,'Helvetica-Bold')
        text(c,260,58,'All four strips start at 0000.',11)
    def closed(c,w,h):
        text(c,20,199,'Front detail - long slider ends are not shown',9)
        x,y=20,35;s=.85
        c.saveState();c.translate(x,y);c.scale(s,s)
        face(c,0,0,True)
        for xx in CENTRES:
            c.setFillColor(SAGE);c.rect(xx-WINDOW_W/2,H/2-WINDOW_H/2,WINDOW_W,WINDOW_H,fill=1,stroke=1)
            digit(c,xx,H/2,'0')
            # Short loop at bottom front, long loop is on the back.
            c.setStrokeColor(HexColor('#555555'));c.setLineWidth(1.4)
            c.roundRect(xx-5,-6,10,23,4,stroke=1,fill=0)
        c.restoreState()
        text(c,420,149,'4 small paper clips',11,'Helvetica-Bold')
        text(c,420,128,'One per slider,',10);text(c,420,113,'on the lower edge.',10)
        text(c,420,84,'Short loop on front;',10);text(c,420,69,'long loop on back.',10)
        text(c,420,36,'Remove to adjust.',10,'Helvetica-Bold')
    def reveal(c,w,h):
        def display(x,chars):
            c.setStrokeColor(GREY);c.rect(x,35,180,66,fill=0,stroke=1)
            for i,ch in enumerate(chars):digit(c,x+27+i*42,68,ch,40)
        display(16,'3705');display(390,'SOLE')
        text(c,20,115,'Upright',10,'Helvetica-Bold');text(c,390,115,'Rotated 180 degrees',10,'Helvetica-Bold')
        c.setStrokeColor(black);c.setLineWidth(1.5);c.line(228,68,350,68)
        p=c.beginPath();p.moveTo(341,74);p.lineTo(350,68);p.lineTo(341,62);c.drawPath(p)
        text(c,245,84,'turn whole tally',10)
        text(c,18,13,'Maker-only check. Do not include this page in the game.',10,'Helvetica-Bold')
    return dict(parts=fig('01_parts',parts),cut=fig('02_cut_windows',cut),
                rails=fig('03_rails',rails),layers=fig('04_layers',layers,height=150),
                install=fig('05_install',install),closed=fig('06_clips',closed),
                reveal=fig('07_maker_check',reveal,height=140))


def doc_base():
    d=Document();s=d.sections[0]
    s.page_width,s.page_height=Inches(8.5),Inches(11)
    s.top_margin=s.bottom_margin=Inches(.55)
    s.left_margin=s.right_margin=Inches(.65)
    normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.05
    for key,size in [('Title',22),('Heading 1',16),('Heading 2',12)]:
        style=d.styles[key];style.font.name='Arial';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(8);style.paragraph_format.space_after=Pt(5)
    # Remove the bundled template's inherited blue title rule.
    for border in list(d.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    d.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    d.core_properties.author='The Cryptic Bureau'
    return d


def print_doc(pdf_path):
    d=doc_base();d.sections[0].top_margin=d.sections[0].bottom_margin=Inches(0)
    d.sections[0].left_margin=d.sections[0].right_margin=Inches(0)
    pdf=pdfium.PdfDocument(str(pdf_path))
    for i in range(len(pdf)):
        if i:d.add_page_break()
        im=pdf[i].render(scale=300/72).to_pil();buf=io.BytesIO();im.save(buf,format='PNG');buf.seek(0)
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(0)
        run=p.add_run();run.add_picture(buf,width=Inches(8.5),height=Inches(11))
        # Page-sized image anchored at 0,0 prevents the inline paragraph adding a page.
        drawing=run._r.find(qn('w:drawing'));inline=drawing.find(qn('wp:inline'));inline.tag=qn('wp:anchor')
        for k,v in {'behindDoc':'0','distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','locked':'0','layoutInCell':'1','allowOverlap':'1','relativeHeight':'1'}.items():inline.set(k,v)
        inline.insert(0,OxmlElement('wp:simplePos'))
        inline[0].set('x','0');inline[0].set('y','0')
        for idx,tag in [(1,'wp:positionH'),(2,'wp:positionV')]:
            pos=OxmlElement(tag);pos.set('relativeFrom','page');off=OxmlElement('wp:posOffset');off.text='0';pos.append(off);inline.insert(idx,pos)
        extent=inline.find(qn('wp:extent'));inline.insert(list(inline).index(extent)+1,OxmlElement('wp:wrapNone'))
    d.core_properties.title='Paper Treasure Tally Print Templates'
    path=OUT/('Print_Paper_Treasure_Tally'+SUFFIX+'.docx');d.save(path);return path


def guide(images):
    d=doc_base();d.core_properties.title='Build the Paper Treasure Tally'
    def p(s):d.add_paragraph(s)
    def heading(s):d.add_paragraph(s,'Heading 2')
    def image(name,width=6.7):
        pp=d.add_paragraph();pp.paragraph_format.space_after=Pt(2)
        pp.add_run().add_picture(str(images[name]),width=Inches(width))
    d.add_paragraph('Build the paper treasure tally','Title')
    if COMPACT:
        p('Envelope-size four-slider prototype. Sleeve: 4 × 2 inches. Number strips: 6 inches long. Stored at 0000, the complete prop is 4 × 6 inches. Digits remain 0.3 inches (7.6 mm) tall. Strips extend farther at mixed settings; keep the tally flat.')
    else:
        p('A reusable four-slider prototype. Sleeve: 6 × 2½ inches. Number strips: 8 inches long. The strips extend well beyond the sleeve during use; keep the tally flat.')
    heading('1  Print at actual size')
    p('Use Letter paper, 100% / actual size, one page per sheet, single-sided. Print template page 1 on tan cardstock and page 2 on light green or white cardstock. Both pages use black ink only. Measure the 1-inch check line before cutting; do not use Fit to page.')
    if COMPACT:
        p('Use both current compact template pages together. The earlier sleeves, rails and wider strips are not interchangeable with this version.')
    heading('You need')
    p('Printer-compatible cardstock, scissors, a craft knife, ruler, cutting mat, glue or narrow double-sided tape, and four small paper clips. Use the cardstock your printer handles reliably.')
    image('parts')
    heading('2  Cut the pieces')
    p('Cut the front, solid backing, ten rails and four complete number strips. Keep both wider tabs on every strip. Cut out the four dashed front windows with a ruler and craft knife. Do not cut the backing guides.')
    image('cut')
    p('Illustrations are assembly diagrams, not cutting templates. Use the separate print-template Word file for all parts.')
    title=d.add_paragraph('Assemble the slider sleeve','Title')
    title.paragraph_format.page_break_before=True
    heading('3  Build five rail stacks')
    p('Glue the ten rail pieces together in pairs. Let them dry flat. Glue one two-layer stack onto each dashed zone on the backing. Keep all four channels completely clear of adhesive.')
    image('rails',6.2)
    image('layers',5.5)
    heading('4  Place the sliders before closing')
    p('With the backing face up, lay one strip in each channel, printed digits facing up and 0 nearest the top. All strips are identical. Keep the wider end tabs outside the sleeve. Do not fold the tabs or glue the strips.')
    image('install',6.0)
    heading('5  Attach the front')
    p('Dry-fit the front over the rails first. Centre 0000 in the windows and check that each strip moves. Then attach the front to the tops of the five rail stacks only. Keep adhesive out of the channels and off the sliders. Leave the top and bottom of all channels open.')
    p('If the strips bind, widen a rough window edge or add one matching rail layer to each of the five stacks. Stock and tape thickness vary, so test the fit before permanently closing the sleeve.')
    title=d.add_paragraph('Check movement and secure the total','Title')
    title.paragraph_format.page_break_before=True
    heading('6  Test all ten digits')
    p('Slide each strip through 0–9. Every digit must centre fully in its window; the wider tabs should stay outside the case. Set the digits by moving the exposed strip ends. Never force a tab into a channel.')
    heading('7  Hold the setting with paper clips')
    p('Once the total is set, put one small paper clip over the lower edge of the sleeve at each slider. Each clip must grip the front, slider and backing together. Put the short loop on the front so it stays below the window, with the long loop on the back. Remove the clips whenever you adjust the digits.')
    image('closed')
    heading('Maker only  Check the puzzle reveal')
    p('This section contains the solution. Keep this guide out of the player envelope. Set 3705, attach all four clips and rotate the entire tally 180 degrees in the plane of the table. Do not flip it onto its back. The digits must stay put and read SOLE.')
    image('reveal')
    heading('Reset and store')
    if COMPACT:
        p('Remove the clips, set 0000 so all strips align, then reattach the clips. The stored prop is 4 × 6 inches, intended for the 6.7 × 4.5 inch envelope. Slide it in with the 6-inch side along the envelope length. Check the actual opening and internal clearance before packing. Do not fold the strips or pack at mixed digit settings.')
    else:
        p('Remove the clips, set 0000, then reattach the clips. With all strips aligned, store the tally flat in a large envelope. Do not crease the number strips.')
    heading('Before play')
    p('The geometry and seven-segment reveal are checked digitally. Cutting tolerances, sliding friction, clip grip and durability still need a real paper test. Check that no neighbouring digit is visible and that the setting survives a full turn.')
    path=OUT/('Paper_Treasure_Tally_Illustrated_Instructions'+SUFFIX+'.docx');d.save(path);return path


def checks():
    assert ''.join(next(ch for ch in 'SOLE' if set(SEGMENTS[ch]) == {ROTATE[s] for s in SEGMENTS[n]}) for n in reversed('3705'))=='SOLE'
    assert abs(STRIP_L-FIRST-9*PITCH-FIRST)<.0001
    assert FIRST-H/2-CAP_H>0  # caps remain outside the case at either extreme
    assert STRIP_W < (RAIL_PITCH-RAIL_W) < CAP_W
    assert DIGIT_H < WINDOW_H < PITCH
    assert (DIGIT_H+WINDOW_H)/2 < PITCH  # neighbouring digits stay outside window
    assert all(RAILS[k]+RAIL_W < CENTRES[k] < RAILS[k+1] for k in range(4))
    assert all(CENTRES[k]-STRIP_W/2 > RAILS[k]+RAIL_W and CENTRES[k]+STRIP_W/2 < RAILS[k+1] for k in range(4))
    assert CENTRES[0]-CAP_W/2 > 0 and CENTRES[-1]+CAP_W/2 < W
    if COMPACT: assert W/I == 4 and STRIP_L/I == 6
    report={'variant':'compact' if COMPACT else 'original',
            'sleeve_inches':[W/I,H/I],'strip_length_inches':STRIP_L/I,'strip_width_inches':STRIP_W/I,
            'stored_at_0000_inches':[W/I,STRIP_L/I],
            'channel_clear_width_inches':round((RAIL_PITCH-RAIL_W)/I,3),'cap_width_inches':CAP_W/I,
            'end_tab_clearance_at_extreme_inches':round((FIRST-H/2-CAP_H)/I,3),
            'digit_height_inches':DIGIT_H/I,'digit_pitch_inches':PITCH/I,'window_inches':[WINDOW_W/I,WINDOW_H/I],
            'rotation_check':'3705 -> SOLE','physical_test':'not performed'}
    (WORK/'geometry_checks.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    checks();pdf=templates();images=figures();print(print_doc(pdf));print(guide(images))
