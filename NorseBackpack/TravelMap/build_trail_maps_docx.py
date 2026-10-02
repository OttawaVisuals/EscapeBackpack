"""Word copy of the four trail maps, one Letter page each: output/docx/Norse_Trail_Maps_Print.docx.

Each map's Print.pdf page 1 is rasterised at 300 dpi, cropped to the sheet's neat line (NEAT in
build_trail_maps_pdf.py), and placed inline at its exact printed size, with page margins equal to
the PDF's own white border. The map therefore prints at the same size and position from Word as
from the PDF -- which matters for Rollo's sheet, measured with the 1:125 ruler (PZ-14). Harald's
second page (the hnefatafl board back) is left out, as in the first hand-made version of this file.

    python NorseBackpack/TravelMap/build_trail_maps_docx.py

Run after build_trail_maps_pdf.py. Written 2 Oct 2026; the 20 Sept version had no script.
"""
import io
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build_trail_maps_pdf import NEAT, PAGE_W, PAGE_H  # noqa: E402

ROOT = HERE.parents[1]
PDF_DIR = ROOT / "output" / "pdf"
OUT = ROOT / "output" / "docx" / "Norse_Trail_Maps_Print.docx"
MAPS = ("Trail_Map_1_Leif_Print", "Trail_Map_2_Rollo_Print",
        "Trail_Map_3_Aud_Print", "Trail_Map_4_Harald_Print")
DPI = 300


def main():
    x0, y0, x1, y1 = NEAT                       # PDF points, origin bottom-left
    clip = pymupdf.Rect(x0, PAGE_H - y1, x1, PAGE_H - y0)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(PAGE_W / 72), Inches(PAGE_H / 72)
    sec.left_margin, sec.right_margin = Inches(x0 / 72), Inches((PAGE_W - x1) / 72)
    sec.top_margin, sec.bottom_margin = Inches((PAGE_H - y1) / 72), Inches(y0 / 72)
    sec.header_distance = sec.footer_distance = Inches(0)
    style = doc.styles["Normal"].paragraph_format
    style.space_before = style.space_after = Pt(0)
    style.line_spacing = 1.0
    for name in MAPS:
        with pymupdf.open(PDF_DIR / (name + ".pdf")) as pdf:
            pix = pdf[0].get_pixmap(matrix=pymupdf.Matrix(DPI / 72, DPI / 72), clip=clip,
                                    alpha=False)
        doc.add_paragraph().add_run().add_picture(
            io.BytesIO(pix.tobytes("png")),
            width=Inches((x1 - x0) / 72), height=Inches((y1 - y0) / 72))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
